#!/usr/bin/env python3
"""把插件 latest.json 汇总成 Excel（openpyxl）。调用方式：
   python3 make_report.py <latest.json> <outDir>
输出：<outDir>/二手车监控-最新.xlsx
"""
import json
import os
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(bold=True, size=13, color="1F3864")
THIN = Side(style="thin", color="D9E1F2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
BAD_FILL = PatternFill("solid", fgColor="FCE4D6")
NEW_FILL = PatternFill("solid", fgColor="E2EFDA")


def wan(v):
    if v is None or v == "":
        return None
    try:
        return round(float(v) / 10000, 2)
    except Exception:
        return None


def load_json(path, default):
    try:
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return default


def load_jsonl(path, limit=200):
    rows = []
    try:
        with open(path, encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if line:
                    rows.append(json.loads(line))
    except Exception:
        pass
    return rows[-limit:]


def add_table(ws, headers, rows, widths, wrap_cols=()):
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    ws.freeze_panes = "A2"
    for row in rows:
        ws.append(row)
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.font = Font(size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=(c in wrap_cols))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main():
    latest_path = sys.argv[1]
    out_dir = sys.argv[2]
    latest = load_json(latest_path, {})
    targets = latest.get("targets", [])
    listings = latest.get("listings", [])
    generated = latest.get("generatedAt", datetime.now().isoformat())
    generated_display = generated[:19].replace("T", " ")
    city = latest.get("city", "西安")

    wb = Workbook()

    # 1. 二手行情
    ws = wb.active
    ws.title = "二手行情"
    headers = ["车型", "样本数", "外市剔除", "页面total", "共找到", "最低(万)", "P25(万)", "中位(万)",
               "P75(万)", "最高(万)", "中位里程(km)", "准新车占比", "参考折价率", "本次新增", "本次降价", "状态", "扫描时间"]
    rows = []
    for t in targets:
        p = t.get("price") or {}
        ref = t.get("reference") or {}
        rate = ref.get("discountVsOfficialLow")
        rows.append([
            t.get("seriesName"), t.get("sampleCount"), t.get("outOfCityCount"), t.get("pageTotal"), t.get("pageFound"),
            wan(p.get("min")), wan(p.get("p25")), wan(p.get("median")), wan(p.get("p75")), wan(p.get("max")),
            (t.get("mileage") or {}).get("medianKm"), (t.get("mileage") or {}).get("nearlyNewRatio"),
            f"{rate * 100:.1f}%" if rate is not None else "—",
            len(t.get("newListings") or []), len(t.get("priceDrops") or []), t.get("status"), generated_display,
        ])
    add_table(ws, headers, rows, [12, 8, 9, 9, 9, 10, 9, 9, 9, 10, 12, 10, 11, 9, 9, 8, 19])
    for r in range(2, ws.max_row + 1):
        status = ws.cell(row=r, column=16).value
        if status == "failed":
            for c in range(1, len(headers) + 1):
                ws.cell(row=r, column=c).fill = BAD_FILL
        elif status == "partial":
            for c in range(1, len(headers) + 1):
                ws.cell(row=r, column=c).fill = WARN_FILL

    # 2. 二手车源
    ws2 = wb.create_sheet("二手车源")
    headers2 = ["车型", "listing_id", "标题", "年款", "配置", "上牌年", "里程(km)", "城市", "挂牌价(万)", "累计降幅(万)", "标签", "详情链接"]
    rows2 = []
    for l in listings:
        rows2.append([
            l.get("seriesName"), l.get("listingId"), l.get("title"), l.get("modelYear"), l.get("trim"),
            l.get("licenseYear"), l.get("mileageKm"), l.get("cityDesc"), wan(l.get("priceYuan")),
            wan(l.get("priceDropYuan")), "、".join(l.get("tags") or []), l.get("detailUrl"),
        ])
    add_table(ws2, headers2, rows2, [12, 22, 40, 8, 34, 8, 10, 8, 11, 12, 18, 44], wrap_cols=(3, 5, 11, 12))

    # 3. 最新上架（本次扫描新增）
    ws3 = wb.create_sheet("最新上架")
    headers3 = ["车型", "标题", "年款", "配置", "上牌年", "里程(km)", "挂牌价(万)", "累计降幅(万)", "标签", "详情链接", "发现时间"]
    rows3 = []
    for t in targets:
        for l in t.get("newListings") or []:
            rows3.append([
                t.get("seriesName"), l.get("title"), l.get("modelYear"), l.get("trim"), l.get("licenseYear"),
                l.get("mileageKm"), wan(l.get("priceYuan")), wan(l.get("priceDropYuan")),
                "、".join(l.get("tags") or []), l.get("detailUrl"), generated_display,
            ])
    add_table(ws3, headers3, rows3, [12, 40, 8, 34, 8, 10, 11, 12, 18, 44, 19], wrap_cols=(2, 4, 9, 10))
    for r in range(2, ws3.max_row + 1):
        for c in range(1, len(headers3) + 1):
            ws3.cell(row=r, column=c).fill = NEW_FILL

    # 4. 降价记录（本次扫描）
    ws4 = wb.create_sheet("降价记录")
    headers4 = ["车型", "标题", "原价(万)", "现价(万)", "降幅(万)", "详情链接", "发现时间"]
    rows4 = []
    for t in targets:
        for d in t.get("priceDrops") or []:
            l = d.get("listing") or {}
            rows4.append([
                t.get("seriesName"), l.get("title"), wan(d.get("oldPrice")), wan(d.get("newPrice")),
                wan((d.get("oldPrice") or 0) - (d.get("newPrice") or 0)), l.get("detailUrl"), generated_display,
            ])
    add_table(ws4, headers4, rows4, [12, 40, 10, 10, 10, 44, 19], wrap_cols=(2, 6))

    # 5. 扫描日志
    ws5 = wb.create_sheet("扫描日志")
    headers5 = ["时间", "触发", "车型数", "成功", "失败", "新增车源", "降价", "推送消息", "报表路径"]
    rows5 = []
    for r in load_jsonl(os.path.join(out_dir, "scan-log.jsonl")):
        rows5.append([
            (r.get("scannedAt") or "")[:19].replace("T", " "), r.get("trigger"), r.get("targets"),
            r.get("okTargets"), r.get("failedTargets"), r.get("newListings"), r.get("priceDrops"),
            r.get("messages"), r.get("report") or "",
        ])
    add_table(ws5, headers5, rows5, [19, 8, 8, 8, 8, 10, 8, 10, 46])

    # 6. 说明
    ws6 = wb.create_sheet("说明")
    ws6["A1"] = f"二手车监控 · {city} · 瓜子二手车"
    ws6["A1"].font = TITLE_FONT
    notes = [
        "监控车型：小米SU7 / 理想i6 / 问界M7（西安）。",
        "数据来源：瓜子二手车车系页（第 1 页，40 条/车型），匿名低频抓取，每 2 小时一次。",
        "口径：挂牌价，不是成交价；只统计西安车源；每页 40 条是样本，所有指标为「样本内」。",
        "新增/降价：与上一次扫描的 listing_id 集合/价格对比；首次扫描只建基线，不逐条推送。",
        "参考折价率 = 1 − 样本中位挂牌价 ÷ 懂车帝新车指导价下限；未做年款/配置匹配，仅看趋势。",
        "页面 total 与「共找到」是两套口径（约 2.25 倍差异），本表分列展示。",
        "合规：仅个人自用、低频、不公开分发、不商用；详情请回源到瓜子页面。",
    ]
    for i, note in enumerate(notes, start=3):
        ws6[f"A{i}"] = note
        ws6[f"A{i}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws6.column_dimensions["A"].width = 110

    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, "二手车监控-最新.xlsx")
    wb.save(out)
    print(f"xlsx: {out}")


if __name__ == "__main__":
    main()
